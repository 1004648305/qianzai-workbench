#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
育儿假 · 初始数据生成脚本
------------------------------------------------------------
用途：把「育儿假.xlsx」转成工作台用的 parental_seed.js。

用法：
    1. 改下面的 SRC（Excel 路径）和 VERSION（版本号，改新即可让页面重新合并）
    2. 运行：python tools/gen_parental_seed.py
    3. 提交并推送 parental_seed.js

合并规则（见 app.js 的 importParentalSeed）：
    - 以「姓名 + 出生日期」作为唯一键
    - 键已存在 → 用表格里的值覆盖（提示栏 note 与「已处理」标记保留，不会被覆盖）
    - 键不存在 → 新增
    - VERSION 不变则跳过合并，因此不会每次刷新都重写

CORRECTIONS：用于修正源表里已知的书写错误（仅作用于输出，不改动用户的 Excel）。
"""

import datetime
import json
import os
import re

from openpyxl import load_workbook

SRC = r"C:\Users\RWX\Desktop\育儿假.xlsx"
OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "parental_seed.js")
VERSION = "20260924b"
SHEET = "Sheet1"

# 字段顺序：与工作台表单一致（表头的中文列名 → 内部字段名）
COLS = [
    ("姓名", "name"),
    ("一级部门", "dept1"),
    ("二级部门", "dept2"),
    ("申请日期", "applyDate"),
    ("出生日期", "birthDate"),
    ("0-1周岁", "p01"),
    ("1-2周岁", "p12"),
    ("2-3周岁", "p23"),
    ("到期日期", "dueDate"),
]

# 已确认的书写错误修正：{姓名: {字段: 修正值}}
CORRECTIONS = {
    "王欢": {"p23": "2026/04/21-2027/04/06"},   # 原 2026/04/21-2027-04/06（分隔符混用）
    "胡琴": {"p12": "2026/08/16-2027/08/15"},   # 原 2026/08/16/-2027/08/15（多一个斜杠）
    "唐伟": {"p12": "2026/08/16-2027/08/15"},   # 同上
}


def norm(v):
    """单元格值归一化：日期→YYYY-MM-DD，整数浮点→整数字符串，其余 strip。"""
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, datetime.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def main():
    wb = load_workbook(SRC, data_only=True)
    ws = wb[SHEET]
    rows = list(ws.iter_rows(values_only=True))

    cn_headers = [norm(c) for c in rows[0]]
    idx = {}
    for cn, field in COLS:
        if cn not in cn_headers:
            raise SystemExit("源表中找不到列：%s（实际表头：%s）" % (cn, cn_headers))
        idx[field] = cn_headers.index(cn)

    records = []
    for r in rows[1:]:
        vals = [norm(c) for c in r]
        if not any(vals):
            continue
        rec = {field: vals[i] for field, i in idx.items()}
        if not rec["name"]:
            continue
        for field, fixed in CORRECTIONS.get(rec["name"], {}).items():
            print("  修正 %s.%s: %s -> %s" % (rec["name"], field, rec[field], fixed))
            rec[field] = fixed
        records.append(rec)

    # 数据质量自检：区间格式
    bad = []
    for rec in records:
        for field in ("p01", "p12", "p23"):
            v = rec[field]
            if v and not re.fullmatch(r"\d{4}/\d{1,2}/\d{1,2}-\d{4}/\d{1,2}/\d{1,2}", v):
                bad.append((rec["name"], field, v))

    fields = [f for _, f in COLS]
    lines = []
    for rec in records:
        lines.append("  " + json.dumps([rec[f] for f in fields], ensure_ascii=False) + ",")

    js = """/* ============================================================
   育儿假 · 初始数据（自动生成，请勿手改）
   来源：{src}
   生成：{ver}  共 {n} 条
   生成脚本：tools/gen_parental_seed.py
   字段顺序：{fields}
   合并规则：按「姓名 + 出生日期」匹配；已存在则更新表格字段，
             保留页面上的「提示栏」与「已处理」标记；不存在则新增。
   ============================================================ */
window.PARENTAL_SEED = {{
  version: '{ver}',
  cols: {cols},
  rows: [
{rows}
  ]
}};
""".format(
        src=SRC,
        ver=VERSION,
        n=len(records),
        fields=" / ".join(cn for cn, _ in COLS),
        cols=json.dumps(fields),
        rows="\n".join(lines),
    )

    with open(OUT, "w", encoding="utf-8") as f:
        f.write(js)

    print("已生成: %s" % OUT)
    print("版本号: %s | 记录数: %d" % (VERSION, len(records)))
    if bad:
        print("\n⚠ 仍有格式异常的区间值（请确认是否已修正）：")
        for name, field, v in bad:
            print("   %s %s = %s" % (name, field, v))
    else:
        print("✓ 区间格式全部规范")


if __name__ == "__main__":
    main()
